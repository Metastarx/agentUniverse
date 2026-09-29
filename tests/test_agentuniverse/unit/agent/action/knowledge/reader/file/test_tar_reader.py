#!/usr/bin/env python3
# -*- coding:utf-8 -*-

# @FileName: test_tar_reader.py
"""Unit tests for the TAR knowledge reader, suffix resolution and registration."""

import importlib.util
import io
import tarfile
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from agentuniverse.agent.action.knowledge.knowledge import Knowledge
from agentuniverse.agent.action.knowledge.reader.file.file_reader import DEFAULT_FILE_READERS, FileReader
from agentuniverse.agent.action.knowledge.reader.file.tar_reader import TarReader
from agentuniverse.agent.action.knowledge.reader.reader_manager import ReaderManager


def _module_available(name: str) -> bool:
    """Return whether an optional module can be imported."""
    return importlib.util.find_spec(name) is not None


def _markdown_available() -> bool:
    """Return whether the markdown sub-reader's optional dependencies exist."""
    return _module_available("unstructured") and _module_available("markdown")


class TestTarReader(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.reader = TarReader()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _add_member(self, archive: tarfile.TarFile, name: str, content) -> None:
        """Append a regular file member to an open TAR archive."""
        if isinstance(content, str):
            content = content.encode("utf-8")
        info = tarfile.TarInfo(name)
        info.size = len(content)
        archive.addfile(info, io.BytesIO(content))

    def _make_tar_bytes(self, name: str, content: str) -> bytes:
        """Return the raw bytes of a single-member TAR archive."""
        buffer = io.BytesIO()
        with tarfile.open(fileobj=buffer, mode="w") as archive:
            self._add_member(archive, name, content)
        return buffer.getvalue()

    def _create_docx_file(self, text: str) -> bytes:
        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph(text)
            buffer = io.BytesIO()
            doc.save(buffer)
            buffer.seek(0)
            return buffer.read()
        except ImportError:
            return b""

    def _create_pptx_file(self, text: str) -> bytes:
        try:
            from pptx import Presentation
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = text
            buffer = io.BytesIO()
            prs.save(buffer)
            buffer.seek(0)
            return buffer.read()
        except ImportError:
            return b""

    def _create_pdf_file(self, text: str) -> bytes:
        try:
            from reportlab.pdfgen import canvas
            from reportlab.lib.pagesizes import letter
            buffer = io.BytesIO()
            pdf = canvas.Canvas(buffer, pagesize=letter)
            pdf.drawString(100, 750, text)
            pdf.save()
            buffer.seek(0)
            return buffer.read()
        except ImportError:
            return b""

    def test_load_plain_tar(self) -> None:
        archive_path = Path(self.temp_dir.name) / "sample.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "docs/readme.txt", "hello world")

        docs = self.reader._load_data(archive_path)

        self.assertEqual(len(docs), 1)
        doc = docs[0]
        self.assertEqual(doc.text, "hello world")
        self.assertEqual(doc.metadata["file_name"], "readme.txt")
        self.assertEqual(doc.metadata["archive_root"], "sample.tar")
        self.assertEqual(doc.metadata["archive_path"], "docs/readme.txt")
        self.assertEqual(doc.metadata["archive_depth"], 0)

    def test_read_compressed_tar_variants(self) -> None:
        variants = (
            ("sample.tar.gz", "w:gz"),
            ("sample.tgz", "w:gz"),
            ("sample.tar.bz2", "w:bz2"),
            ("sample.tar.xz", "w:xz"),
        )
        for file_name, mode in variants:
            archive_path = Path(self.temp_dir.name) / file_name
            with tarfile.open(archive_path, mode) as archive:
                self._add_member(archive, "docs/readme.txt", "hello world")

            docs = self.reader._load_data(archive_path)

            self.assertEqual(len(docs), 1, file_name)
            self.assertEqual(docs[0].text, "hello world")
            self.assertEqual(docs[0].metadata["file_name"], "readme.txt")

    def test_multiple_text_like_members(self) -> None:
        archive_path = Path(self.temp_dir.name) / "mixed.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "notes.txt", "plain text")
            self._add_member(archive, "data.csv", "name,age\nalice,30")
            self._add_member(archive, "config.json", '{"debug": true}')
            self._add_member(archive, "script.py", "print('hi')")
            self._add_member(archive, "service.conf", "workers=4")
            self._add_member(archive, "app.log", "started")

        docs = self.reader._load_data(archive_path)

        names = {doc.metadata["file_name"] for doc in docs}
        self.assertEqual(len(docs), 6)
        self.assertEqual(
            names,
            {"notes.txt", "data.csv", "config.json", "script.py", "service.conf", "app.log"},
        )

    def test_markdown_member_is_routed(self) -> None:
        if not _markdown_available():
            self.skipTest("unstructured/markdown are not installed")

        archive_path = Path(self.temp_dir.name) / "docs.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "README.md", "# Title\n\nbody")

        docs = self.reader._load_data(archive_path)

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["file_name"], "README.md")

    def test_optional_office_members(self) -> None:
        members = {}
        if _module_available("docx") and _module_available("docx2txt"):
            members["report.docx"] = self._create_docx_file("word content")
        if _module_available("pptx"):
            members["slides.pptx"] = self._create_pptx_file("slide title")
        if _module_available("reportlab") and _module_available("pypdf"):
            members["document.pdf"] = self._create_pdf_file("pdf content")

        if not members:
            self.skipTest("optional docx/pptx/pdf dependencies are not installed")

        archive_path = Path(self.temp_dir.name) / "office.tar"
        with tarfile.open(archive_path, "w") as archive:
            for name, content in members.items():
                self._add_member(archive, name, content)

        docs = self.reader._load_data(archive_path)
        names = {doc.metadata.get("file_name") for doc in docs}

        for name in members:
            self.assertIn(name, names)

    def test_nested_tar_is_expanded(self) -> None:
        inner = self._make_tar_bytes("inner/data.txt", "nested data")
        archive_path = Path(self.temp_dir.name) / "outer.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "folder/archive.tar", inner)

        docs = self.reader._load_data(archive_path)

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].text, "nested data")
        self.assertEqual(docs[0].metadata["archive_path"], "folder/archive.tar/inner/data.txt")
        self.assertEqual(docs[0].metadata["archive_depth"], 1)

    def test_nested_tar_depth_limit(self) -> None:
        current = self._make_tar_bytes("data.txt", "deepest")
        for index in range(10):
            buffer = io.BytesIO()
            with tarfile.open(fileobj=buffer, mode="w") as archive:
                self._add_member(archive, f"level{index}.tar", current)
            current = buffer.getvalue()

        archive_path = Path(self.temp_dir.name) / "deep.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "nested.tar", current)

        shallow_reader = TarReader(max_depth=2)
        with self.assertRaises(ValueError) as context:
            shallow_reader._load_data(archive_path)
        self.assertIn("nesting depth", str(context.exception))

    def test_file_count_limit(self) -> None:
        archive_path = Path(self.temp_dir.name) / "many.tar"
        with tarfile.open(archive_path, "w") as archive:
            for index in range(100):
                self._add_member(archive, f"file_{index}.txt", f"content {index}")

        limited_reader = TarReader(max_files=50)
        with self.assertRaises(ValueError) as context:
            limited_reader._load_data(archive_path)
        self.assertIn("maximum file count", str(context.exception))

    def test_total_size_limit(self) -> None:
        archive_path = Path(self.temp_dir.name) / "large_total.tar"
        with tarfile.open(archive_path, "w") as archive:
            for index in range(20):
                self._add_member(archive, f"file_{index}.txt", "x" * 1000)

        limited_reader = TarReader(max_total_size=5000)
        with self.assertRaises(ValueError) as context:
            limited_reader._load_data(archive_path)
        self.assertIn("maximum total size", str(context.exception))

    def test_exceeds_file_size_limit(self) -> None:
        archive_path = Path(self.temp_dir.name) / "limit.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "large.txt", "a" * 4096)

        limited_reader = TarReader(max_file_size=1024, max_total_size=2048)
        with self.assertRaises(ValueError):
            limited_reader._load_data(archive_path)

    def test_compression_ratio_limit(self) -> None:
        archive_path = Path(self.temp_dir.name) / "bomb.tar.gz"
        with tarfile.open(archive_path, "w:gz") as archive:
            self._add_member(archive, "repetitive.txt", "a" * 100000)

        strict_reader = TarReader(max_compression_ratio=10)
        with self.assertRaises(ValueError) as context:
            strict_reader._load_data(archive_path)
        self.assertIn("compression ratio", str(context.exception))

    def test_path_traversal_members_are_skipped(self) -> None:
        archive_path = Path(self.temp_dir.name) / "traversal.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "../../../etc/passwd", "blocked")
            self._add_member(archive, "/absolute.txt", "blocked")
            self._add_member(archive, "./../sensitive.txt", "blocked")
            self._add_member(archive, "normal/file.txt", "allowed")

        docs = self.reader._load_data(archive_path)

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["archive_path"], "normal/file.txt")
        for doc in docs:
            self.assertNotIn("..", doc.metadata.get("archive_path", ""))

    def test_symlink_members_are_skipped(self) -> None:
        archive_path = Path(self.temp_dir.name) / "links.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "real.txt", "real content")
            link = tarfile.TarInfo("link.txt")
            link.type = tarfile.SYMTYPE
            link.linkname = "real.txt"
            link.size = 0
            archive.addfile(link)

        docs = self.reader._load_data(archive_path)

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["file_name"], "real.txt")

    def test_custom_metadata_is_propagated(self) -> None:
        archive_path = Path(self.temp_dir.name) / "meta.tar"
        with tarfile.open(archive_path, "w") as archive:
            self._add_member(archive, "docs/file.txt", "content")

        docs = self.reader._load_data(archive_path, ext_info={"project": "demo", "version": "1.0"})

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].metadata["project"], "demo")
        self.assertEqual(docs[0].metadata["version"], "1.0")
        self.assertEqual(docs[0].metadata["file_name"], "file.txt")

    def test_missing_file_raises(self) -> None:
        with self.assertRaises(FileNotFoundError):
            self.reader._load_data(Path(self.temp_dir.name) / "missing.tar")

    def test_non_path_input_raises(self) -> None:
        with self.assertRaises(TypeError):
            self.reader._load_data(123)


class TestFileReaderCompoundSuffix(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_resolve_suffix_prefers_compound(self) -> None:
        file_reader = FileReader()

        self.assertEqual(file_reader._resolve_suffix(Path("bundle.tar.gz")), ".tar.gz")
        self.assertEqual(file_reader._resolve_suffix(Path("plain.tgz")), ".tgz")
        self.assertEqual(file_reader._resolve_suffix(Path("plain.tar")), ".tar")
        self.assertEqual(file_reader._resolve_suffix(Path("notes.txt")), ".txt")
        self.assertEqual(file_reader._resolve_suffix(Path("data.extra.txt")), ".txt")
        self.assertEqual(file_reader._resolve_suffix(Path("noextension")), "")

    def test_load_data_dispatches_compound_suffix(self) -> None:
        archive_path = Path(self.temp_dir.name) / "bundle.tar.gz"
        with tarfile.open(archive_path, "w:gz") as archive:
            info = tarfile.TarInfo("notes.txt")
            payload = b"hello from tar.gz"
            info.size = len(payload)
            archive.addfile(info, io.BytesIO(payload))

        docs = FileReader()._load_data([archive_path])

        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0].text, "hello from tar.gz")
        self.assertEqual(docs[0].metadata["file_name"], "notes.txt")


class TestKnowledgeCompoundSuffix(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_load_data_resolves_tar_gz(self) -> None:
        archive_path = Path(self.temp_dir.name) / "bundle.tar.gz"
        with tarfile.open(archive_path, "w:gz") as archive:
            info = tarfile.TarInfo("notes.txt")
            payload = b"hello"
            info.size = len(payload)
            archive.addfile(info, io.BytesIO(payload))

        knowledge = Knowledge(name="tar_suffix_test")
        stub_reader = mock.MagicMock()
        stub_reader.load_data.return_value = []

        with mock.patch("agentuniverse.agent.action.knowledge.knowledge.ReaderManager") as manager_cls:
            manager_cls.return_value.get_file_default_reader.return_value = stub_reader
            knowledge._load_data(source_path=str(archive_path))
            manager_cls.return_value.get_file_default_reader.assert_called_once_with("tar.gz")


class TestReaderRegistration(unittest.TestCase):
    def test_default_reader_registration(self) -> None:
        default_reader = ReaderManager().DEFAULT_READER

        self.assertEqual(default_reader.get("tar"), "default_tar_reader")
        self.assertEqual(default_reader.get("tgz"), "default_tar_reader")
        self.assertEqual(default_reader.get("tar.gz"), "default_tar_reader")

    def test_default_file_readers_registration(self) -> None:
        for suffix in (".tar", ".tgz", ".tar.gz", ".tar.bz2", ".tar.xz"):
            self.assertIs(DEFAULT_FILE_READERS[suffix], TarReader)


if __name__ == "__main__":
    unittest.main()
